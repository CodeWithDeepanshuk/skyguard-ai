from pathlib import Path
import openpyxl
from fastapi.testclient import TestClient
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT))

from skyguard.api.app import create_app


def test_excel_export():
    app = create_app(root=ROOT)
    client = TestClient(app)
    res = client.get("/api/export/weather_anomalies.xlsx")
    assert res.status_code == 200, f"Expected 200, got {res.status_code}"
    assert "openxmlformats" in res.headers.get("content-type", "")
    assert len(res.content) > 5000, f"Excel file size too small: {len(res.content)}"

    # Load into openpyxl from bytes
    import io
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Weather Anomaly Analysis" in wb.sheetnames
    assert "Summary & Statistics" in wb.sheetnames
    assert "Fault Guide" in wb.sheetnames

    ws = wb["Weather Anomaly Analysis"]
    assert ws.max_row >= 31, f"Expected >= 31 rows, got {ws.max_row}"
    # Verify header red text
    assert ws["A1"].font.color.rgb == "00C00000"
    print("SUCCESS: Excel export endpoint returns valid styled workbook with sheets:", wb.sheetnames)


if __name__ == "__main__":
    test_excel_export()
